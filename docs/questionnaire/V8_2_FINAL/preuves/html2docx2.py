"""HTML simple (h1, h2, h3, p, ul/li, b, i, br, table/tr/th/td) → .docx."""
import sys, re
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

class Conv(HTMLParser):
    def __init__(self, doc):
        super().__init__(convert_charrefs=True)
        self.doc = doc; self.p = None; self.bold = 0; self.ital = 0; self.list_depth = 0
        self.color = None; self.skip = False; self.table = None; self.row = None; self.cell = None; self.th = False
    def new_par(self, style=None, attrs=None):
        self.p = self.doc.add_paragraph(style=style); self.color = None
        st = dict(attrs or {}).get('style', '') or ''
        m = re.search(r'margin-left:\s*(\d+)pt', st)
        if m: self.p.paragraph_format.left_indent = Pt(int(m.group(1)))
        c = re.search(r'color:\s*#([0-9A-Fa-f]{6})', st)
        if c: self.color = RGBColor.from_string(c.group(1).upper())
    def handle_starttag(self, tag, attrs):
        if tag in ('title', 'head'): self.skip = True
        elif tag == 'table': self.table = []
        elif tag == 'tr': self.row = []
        elif tag in ('td', 'th'): self.cell = []; self.th = tag == 'th'
        elif self.cell is not None:
            if tag == 'b': self.bold += 1
            elif tag == 'i': self.ital += 1
            elif tag == 'br': self.cell.append(('\n', False, False))
        elif tag == 'h1': self.p = self.doc.add_heading(level=0); self.color = None
        elif tag == 'h2': self.p = self.doc.add_heading(level=1); self.color = None
        elif tag == 'h3': self.p = self.doc.add_heading(level=2); self.color = None
        elif tag == 'p': self.new_par(attrs=attrs)
        elif tag == 'ul': self.list_depth += 1
        elif tag == 'li': self.new_par(style='List Bullet' if self.list_depth <= 1 else 'List Bullet 2')
        elif tag == 'b': self.bold += 1
        elif tag == 'i': self.ital += 1
        elif tag == 'br' and self.p is not None: self.p.add_run().add_break(WD_BREAK.LINE)
    def handle_endtag(self, tag):
        if tag in ('title', 'head'): self.skip = False
        elif tag in ('td', 'th'): self.row.append((self.cell, self.th)); self.cell = None
        elif tag == 'tr': self.table.append(self.row); self.row = None
        elif tag == 'table': self.flush_table(); self.table = None
        elif tag in ('p', 'h1', 'h2', 'h3', 'li'): self.p = None
        elif tag == 'ul': self.list_depth -= 1
        elif tag == 'b': self.bold -= 1
        elif tag == 'i': self.ital -= 1
    def flush_table(self):
        rows = self.table; ncol = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncol); t.style = 'Table Grid'
        for i, r in enumerate(rows):
            for j, (runs, th) in enumerate(r):
                c = t.cell(i, j); par = c.paragraphs[0]
                for text, b, it in runs:
                    if text == '\n': par.add_run().add_break(WD_BREAK.LINE); continue
                    rr = par.add_run(text); rr.font.size = Pt(8.5); rr.bold = (b or th) or None; rr.italic = it or None
                if th: shade(c, 'EDE7F6')
        self.doc.add_paragraph()
    def handle_data(self, data):
        if self.skip: return
        text = data.replace('\n', ' ')
        if self.cell is not None:
            if text.strip() or text == ' ': self.cell.append((text, self.bold > 0, self.ital > 0))
            return
        if self.table is not None: return
        if not text.strip():
            if self.p is not None and text: self.p.add_run(' ')
            return
        if self.p is None: self.new_par()
        r = self.p.add_run(text); r.bold = self.bold > 0 or None; r.italic = self.ital > 0 or None
        if self.color is not None: r.font.color.rgb = self.color

src, dst = sys.argv[1], sys.argv[2]
doc = Document()
doc.styles['Normal'].font.name = 'Arial'; doc.styles['Normal'].font.size = Pt(10.5)
for s in doc.sections:
    s.left_margin = s.right_margin = Pt(54)
Conv(doc).feed(open(src, encoding='utf-8').read())
doc.save(dst); print(dst)
